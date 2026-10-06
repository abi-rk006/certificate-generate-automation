import sys
import os
import re
from pptx import Presentation

def get_default_paths():
    """
    Determine default paths for template and output directory,
    supporting both local environment (e.g. Windows) and Docker/n8n (/files).
    """
    # 1. Script & Project directories
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)

    # Check for template
    local_template = os.path.join(project_root, "template", "Certificate_template.pptx")
    docker_template = "/files/template/Certificate_template.pptx"

    if os.path.exists(local_template):
        template_path = local_template
    elif os.path.exists(docker_template):
        template_path = docker_template
    else:
        # Fallback to local
        template_path = local_template

    # Check for output directory
    if os.path.exists("/files"):
        output_dir = "/files/output/pptx"
    else:
        output_dir = os.path.join(project_root, "output", "pptx")

    return template_path, output_dir

def sanitize_filename(name: str) -> str:
    """Remove characters that are invalid in filenames."""
    return re.sub(r'[\\/*?:"<>|]', "", name).strip()

def replace_placeholder_in_shape(shape, placeholder: str, replacement: str) -> bool:
    """
    Replace placeholder in shape text while preserving font style, size,
    color, boldness, and alignment.
    """
    if not shape.has_text_frame:
        return False

    replaced = False
    for paragraph in shape.text_frame.paragraphs:
        if placeholder in paragraph.text:
            # 1. Check if the placeholder is contained entirely within a single run
            found_in_single_run = False
            for run in paragraph.runs:
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, replacement)
                    found_in_single_run = True
                    replaced = True

            # 2. If the placeholder spans across multiple runs, preserve run 0 formatting
            if not found_in_single_run:
                new_text = paragraph.text.replace(placeholder, replacement)
                if paragraph.runs:
                    first_run = paragraph.runs[0]
                    font_name = first_run.font.name
                    font_size = first_run.font.size
                    font_bold = first_run.font.bold
                    font_italic = first_run.font.italic
                    font_color = (
                        first_run.font.color.rgb
                        if (first_run.font.color and hasattr(first_run.font.color, "rgb"))
                        else None
                    )

                    first_run.text = new_text
                    for extra_run in paragraph.runs[1:]:
                        extra_run.text = ""

                    if font_name:
                        first_run.font.name = font_name
                    if font_size:
                        first_run.font.size = font_size
                    if font_bold is not None:
                        first_run.font.bold = font_bold
                    if font_italic is not None:
                        first_run.font.italic = font_italic
                    if font_color:
                        first_run.font.color.rgb = font_color
                else:
                    paragraph.text = new_text
                replaced = True

    return replaced

def create_certificate(name: str, template_path: str = None, output_dir: str = None) -> str:
    """
    Generate a certificate PPTX for a given name.
    """
    default_template, default_output = get_default_paths()
    template_path = template_path or default_template
    output_dir = output_dir or default_output

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Certificate template not found at: {template_path}")

    os.makedirs(output_dir, exist_ok=True)

    prs = Presentation(template_path)

    placeholder_found = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and "{{NAME}}" in shape.text_frame.text:
                if replace_placeholder_in_shape(shape, "{{NAME}}", name):
                    placeholder_found = True

    safe_name = sanitize_filename(name) or "certificate"
    output_file = os.path.join(output_dir, f"{safe_name}.pptx")
    prs.save(output_file)

    if not placeholder_found:
        print(f"Warning: '{{{{NAME}}}}' placeholder not found in template for '{name}'. File saved anyway.")

    return output_file

def main():
    if len(sys.argv) < 2:
        print("Usage:")
        print("  1. Single name:     python generate_certificate.py \"Student Name\" [--pdf]")
        print("  2. Multiple names:   python generate_certificate.py \"Name 1\" \"Name 2\" [--pdf]")
        print("  3. Comma-separated:  python generate_certificate.py --names \"Name1, Name2\" [--pdf]")
        sys.exit(1)

    pdf_flag = "--pdf" in sys.argv
    clean_args = [arg for arg in sys.argv[1:] if arg != "--pdf"]

    names = []
    if clean_args and clean_args[0] == "--names" and len(clean_args) > 1:
        names = [n.strip() for n in clean_args[1].split(",") if n.strip()]
    else:
        names = [arg.strip() for arg in clean_args if arg.strip()]

    print(f"Generating certificates for {len(names)} recipient(s)...")
    for idx, name in enumerate(names, 1):
        try:
            output_path = create_certificate(name)
            msg = f"[{idx}/{len(names)}] Created: {output_path}"
            if pdf_flag:
                try:
                    from convert_to_pdf import convert_pptx_to_pdf
                    pdf_path = convert_pptx_to_pdf(output_path)
                    msg += f" & PDF: {pdf_path}"
                except Exception as pe:
                    msg += f" (PDF conversion failed: {pe})"
            print(msg)
        except Exception as e:
            print(f"[{idx}/{len(names)}] Error generating for '{name}': {e}")

if __name__ == "__main__":
    main()
